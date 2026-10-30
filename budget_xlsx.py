"""Builds the Excel version of a budget (same layout as the Art of Us workbook): key variables, then
Quantity / Unit cost / Times / Budget / Notes per line, with live formulas so it recalculates in Excel."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import budget_model as bm

A = Font(name="Arial", size=10)
B = Font(name="Arial", size=10, bold=True)
H = Font(name="Arial", size=14, bold=True)
BL = Font(name="Arial", size=10, color="0000FF")
WH = Font(name="Arial", size=10, bold=True, color="FFFFFF")
YEL = PatternFill("solid", fgColor="FFFF99")
TEAL = PatternFill("solid", fgColor="0C4561")
LT = PatternFill("solid", fgColor="E9F4F7")
TOP = Border(top=Side(style="thin", color="000000"))
M = "$#,##0;($#,##0);-"
M2 = "$#,##0.00;($#,##0.00);-"


def build_xlsx(label, kind, data, updated_at=None):
    """Returns the .xlsx file as bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Discovery!"
    ws["A1"].font = H
    ws["A2"] = f"Budget: {label}"
    ws["A2"].font = B
    ws["A3"] = ("Blue cells on yellow are inputs; everything else calculates. Prices marked ESTIMATE or PLACEHOLDER "
                "are not quotes." + (f" Saved {updated_at}." if updated_at else ""))
    ws["A3"].font = A
    r = 5

    def hdr(cols):
        nonlocal r
        for i, t in enumerate(cols, 1):
            c = ws.cell(r, i, t)
            c.font = WH
            c.fill = TEAL
            c.alignment = Alignment(horizontal="left" if i in (1, 7) else "right", wrap_text=True)
        r += 1

    def inp(row, col, val, fmt):
        c = ws.cell(row, col, val)
        c.font = BL
        c.fill = YEL
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")

    K = {}
    if kind == "training":
        hdr(["KEY VARIABLES", "Value", "", "", "", "", "Note"])
        for k, lab, _d, vkind, note in bm.KEY_VARS:
            ws.cell(r, 1, lab).font = A
            fmt = {"pct": "0.0%", "money2": M2}.get(vkind, "0")
            inp(r, 2, data["vars"].get(k, _d), fmt)
            ws.cell(r, 7, note).font = A
            K[k] = f"$B${r}"
            r += 1
        r += 1
    hdr(["ITEM", "Quantity", "Unit cost", "Times (nights / days)", "", "Budget", "Basis / notes"])

    rows, totals = bm.compute(data)
    fee_row = None
    sec_total_row = {}
    for sec, sec_label in bm.sections_for(kind):
        for c in range(1, 8):
            ws.cell(r, c).fill = LT
        ws.cell(r, 1, sec_label).font = B
        r += 1
        first = r
        for ln in [x for x in data["lines"] if x["section"] == sec]:
            ws.cell(r, 1, "  " + ln["label"]).font = A

            def cell(col, raw, fmt):
                if isinstance(raw, str) and raw.startswith("="):
                    people = f"(2*B{fee_row}+{K['team_people']})" if fee_row and "team_people" in K else None
                    f = {"=couples": f"=B{fee_row}" if fee_row else "=0",
                         "=trainees": f"=2*B{fee_row}" if fee_row else "=0",
                         "=people": f"={people}" if people else "=0",
                         "=rack": f"={K['rack']}" if "rack" in K else "=0"}[raw]
                    c = ws.cell(r, col, f)
                    c.font = A
                    c.number_format = fmt
                    c.alignment = Alignment(horizontal="right")
                else:
                    inp(r, col, raw, fmt)

            cell(2, ln["q"], "#,##0")
            cell(3, ln["u"], M2 if (isinstance(ln["u"], (int, float)) and ln["u"] % 1) else M)
            cell(4, ln["t"], "0")
            f = f"=B{r}*C{r}*D{r}"
            mode = ln.get("mode", "plain")
            if K:
                room = f"*(1-{K['room_discount']})*(1+{K['room_tax']})"
                food = f"*(1-{K['food_discount']})*(1+{K['food_service']})"
                if mode == "room":
                    f += room
                elif mode == "food":
                    f += food
                elif mode == "coffee":
                    f += f"*{K['coffee_on']}" + food
                elif mode == "fee" and fee_row:
                    f += f"+F{fee_row}*{K['card_rate']}"
            c = ws.cell(r, 6, f)
            c.font = A
            c.number_format = M
            ws.cell(r, 7, ln.get("note", "")).font = A
            if ln["id"] == "fee_income":
                fee_row = r
            r += 1
        ws.cell(r, 1, "Total for " + sec_label.split(" (")[0]).font = B
        c = ws.cell(r, 6, f"=SUM(F{first}:F{r - 1})" if r > first else 0)
        c.font = B
        c.number_format = M
        c.border = TOP
        sec_total_row[sec] = r
        r += 2

    def put(label_, formula, fmt=M, bold=False):
        nonlocal r
        ws.cell(r, 1, label_).font = B if bold else A
        c = ws.cell(r, 6, formula)
        c.font = B if bold else A
        c.number_format = fmt
        r += 1
        return r - 1

    S = sec_total_row
    if kind == "training":
        op = f"(F{S['venue']}+F{S['team']}+F{S['meals']}+F{S['other']})"
        cont = put("Contingency", f"={op}*{K['contingency']}")
        te = put("Total expenses (training)", f"={op}+F{cont}", bold=True)
        ni = put("Net for this training (before fixed and one-time expenses)", f"=F{S['income']}-F{te}", bold=True)
        r += 1
        put("Expense per couple", f"=F{te}/B{fee_row}")
        put("Net per couple", f"=F{ni}/B{fee_row}")
        put("Break-even number of couples", f"=ROUNDUP(F{te}/(C{fee_row}*(1-{K['card_rate']})),0)", "0")
    elif kind == "donations":
        te = put("Total expenses", f"=F{S['other']}", bold=True)
        put("Net donations", f"=F{S['income']}-F{te}", bold=True)
    for col, w in zip("ABCDEFG", (52, 12, 13, 13, 3, 15, 95)):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
